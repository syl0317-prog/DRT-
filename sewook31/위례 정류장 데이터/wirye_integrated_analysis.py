
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

# 조사 대상 정류장 13개
stations = [
    "28370", "28469", "28470", "28471", "28473",
    "05174", "05290", "05291", "05292", "05293",
    "05294", "28372", "28468"
]

routes = ["31", "38", "60", "200", "210", "231", "315", "331", "위례-01"]

# 기존 조사 결과: 정류장별 경유 노선
station_routes = {
    "28370": ["31", "38", "위례-01"],
    "28469": ["31", "38", "200", "210", "231", "315", "331", "위례-01"],
    "28470": ["31", "38", "위례-01"],
    "28471": ["31", "38", "위례-01"],
    "28473": ["31", "38", "200", "210", "231", "315", "331", "위례-01"],
    "05174": ["31", "200", "210", "231", "315", "331"],
    "05290": ["31", "200", "210", "231", "315", "331"],
    "05291": ["31", "200", "210", "231", "315", "331"],
    "05292": ["31", "200", "210", "231", "315", "331"],
    "05293": ["31", "200", "210", "231", "315", "331"],
    "05294": ["31", "200", "210", "231", "315", "331"],
    "28372": ["31", "38", "200", "210", "231", "315", "331", "위례-01"],
    "28468": ["31", "38", "200", "210", "231", "315", "331", "위례-01"],
}

# 노선별 배차간격
# 순서: ID, 평일 출퇴근, 평일 일반, 주말 출퇴근, 주말 일반
headways = {
    "31": ["227000017", 20, 15, 20, 30],
    "38": ["227000004", 18, 25, 18, 30],
    "200": ["204000029", 45, 40, 55, 65],
    "210": ["204000180", 25, 15, 25, 35],
    "231": ["204000139", 32, 28, 32, 42],
    "315": ["204000073", 20, 24, 25, 35],
    "331": ["204000084", 25, 28, 25, 35],
    "위례-01": ["미확인", "15~35", "15~35", "23~35", "23~35"],
}

wb = Workbook()

# 시트 1: 정류장별 경유노선
ws1 = wb.active
ws1.title = "정류장별 경유노선"
ws1.append(["정류장번호"] + routes + ["경유노선수"])

for station in stations:
    served = station_routes[station]
    ws1.append(
        [station] +
        ["O" if route in served else "X" for route in routes] +
        [len(served)]
    )

# 시트 2: 노선별 배차간격
ws2 = wb.create_sheet("노선별 배차간격")
ws2.append([
    "버스번호", "노선ID", "평일 출퇴근(분)",
    "평일 일반(분)", "주말 출퇴근(분)", "주말 일반(분)"
])

for route in routes:
    if route in headways:
        ws2.append([route] + headways[route])
    else:
        ws2.append([route, "분석 제외", None, None, None, None])

# 시트 3: 정류장별 배차정보
ws3 = wb.create_sheet("정류장별 배차정보")
ws3.append([
    "정류장번호", "버스번호", "노선ID",
    "평일 출퇴근(분)", "평일 일반(분)",
    "주말 출퇴근(분)", "주말 일반(분)"
])

for station in stations:
    for route in routes:
        if route in station_routes[station] and route in headways:
            ws3.append([station, route] + headways[route])

# 엑셀 서식
for ws in wb.worksheets:
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="EAF1F8")

    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        max_length = max(len(str(c.value or "")) for c in column)
        ws.column_dimensions[letter].width = max(14, min(max_length + 4, 25))

    # 정류장번호 앞의 0 유지
    if ws.title != "노선별 배차간격":
        for row in ws.iter_rows(min_row=2):
            row[0].number_format = "@"

output = Path(__file__).parent / "위례동_13개정류장_버스노선_통합분석.xlsx"
wb.save(output)

print("통합 엑셀 파일 생성 완료!")
print("저장 위치:", output)
print("정류장 수:", len(stations))
print("배차간격 확보 노선:", len(headways))
